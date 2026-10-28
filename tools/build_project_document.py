from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PROJECT_DOCUMENTATION.md"
OUTPUT_DIR = ROOT / "output" / "docx"
OUTPUT = OUTPUT_DIR / "Anaplan_Daily_Backup_Project_Documentation.docx"
FLOWCHART = ROOT / "tmp" / "doc_build" / "backup_process_flow.png"

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EDF3F8"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"
TEXT = RGBColor(31, 31, 31)


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "",
        "/Library/Fonts/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    if bold:
        candidates.insert(0, "/System/Library/Fonts/Supplemental/Arial Bold.ttf")
        candidates.insert(1, "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def rounded_box(draw, xy, text, fill, outline=NAVY, text_fill="FFFFFF"):
    draw.rounded_rectangle(xy, radius=22, fill=f"#{fill}", outline=f"#{outline}", width=4)
    box = draw.multiline_textbbox((0, 0), text, font=font(31, True), spacing=8, align="center")
    width = box[2] - box[0]
    height = box[3] - box[1]
    x1, y1, x2, y2 = xy
    draw.multiline_text(
        ((x1 + x2 - width) / 2, (y1 + y2 - height) / 2 - 2),
        text,
        font=font(31, True),
        fill=f"#{text_fill}",
        spacing=8,
        align="center",
    )


def arrow(draw, start, end, label=None):
    draw.line([start, end], fill=f"#{NAVY}", width=6)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) > abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 18 * direction, y2 - 12), (x2 - 18 * direction, y2 + 12)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 12, y2 - 18 * direction), (x2 + 12, y2 - 18 * direction)]
    draw.polygon(points, fill=f"#{NAVY}")
    if label:
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        draw.text((mx + 10, my - 28), label, font=font(24, True), fill=f"#{NAVY}")


def create_flowchart():
    image = Image.new("RGB", (1700, 1450), "white")
    draw = ImageDraw.Draw(image)

    boxes = {
        "config": (500, 40, 1200, 155),
        "mode": (500, 215, 1200, 330),
        "dry": (1250, 215, 1650, 330),
        "login": (500, 390, 1200, 505),
        "target": (500, 565, 1200, 680),
        "copy": (500, 740, 1200, 855),
        "verify": (500, 915, 1200, 1030),
        "record": (500, 1090, 1200, 1205),
        "more": (500, 1265, 1200, 1380),
        "finish": (50, 1265, 450, 1380),
    }

    rounded_box(draw, boxes["config"], "Load configuration\nand enabled models", LIGHT_BLUE, text_fill=NAVY)
    rounded_box(draw, boxes["mode"], "Validate configuration\nand choose run mode", BLUE)
    rounded_box(draw, boxes["dry"], "Write validation or\ndry-run report", LIGHT_BLUE, text_fill=NAVY)
    rounded_box(draw, boxes["login"], "Acquire run lock, open\nbrowser, authenticate", BLUE)
    rounded_box(draw, boxes["target"], "Open exact workspace\nand model by IDs", BLUE)
    rounded_box(draw, boxes["copy"], "Preflight access or\nCopy and Archive", BLUE)
    rounded_box(draw, boxes["verify"], "For live runs, confirm OK\nand poll for archive", BLUE)
    rounded_box(draw, boxes["record"], "Record success\nor failure", LIGHT_BLUE, text_fill=NAVY)
    rounded_box(draw, boxes["more"], "More enabled\nmodels", LIGHT_BLUE, text_fill=NAVY)
    rounded_box(draw, boxes["finish"], "Close browser and\nwrite final report", NAVY)

    arrow(draw, (850, 155), (850, 215))
    arrow(draw, (1200, 272), (1250, 272))
    arrow(draw, (850, 330), (850, 390), "Live")
    arrow(draw, (850, 505), (850, 565))
    arrow(draw, (850, 680), (850, 740))
    arrow(draw, (850, 855), (850, 915))
    arrow(draw, (850, 1030), (850, 1090))
    arrow(draw, (850, 1205), (850, 1265))
    arrow(draw, (500, 1322), (450, 1322), "No")
    draw.line([(1200, 1322), (1510, 1322), (1510, 622), (1200, 622)], fill=f"#{NAVY}", width=6)
    draw.polygon([(1200, 622), (1220, 610), (1220, 634)], fill=f"#{NAVY}")
    draw.text((1220, 1280), "Yes", font=font(24, True), fill=f"#{NAVY}")

    image.save(FLOWCHART, dpi=(180, 180))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Arial", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def clean_inline(text):
    return re.sub(r"[`*]", "", text).strip()


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8, color="666666")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(16)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 17, 16, 7),
        ("Heading 2", 12.5, 11, 5),
        ("Heading 3", 10.5, 8, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_table(doc, lines):
    rows = []
    for line in lines:
        if re.match(r"^\|(?:\s*:?-+:?\s*\|)+$", line):
            continue
        rows.append([clean_inline(cell) for cell in line.strip().strip("|").split("|")])
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.text = values[column_index] if column_index < len(values) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            set_cell_shading(cell, NAVY if row_index == 0 else (PALE_BLUE if row_index % 2 == 0 else "FFFFFF"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=8.2, bold=row_index == 0, color="FFFFFF" if row_index == 0 else "1F1F1F")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    row_properties = table.rows[0]._tr.get_or_add_trPr()
    row_properties.append(OxmlElement("w:cantSplit"))
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_borders(cell, "C8C8C8")
    set_cell_margins(cell, 130, 150, 130, 150)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("\n".join(lines))
    set_run_font(run, name="Courier New", size=7.7, color="202020")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_markdown_content(doc, markdown):
    lines = markdown.splitlines()
    index = 1
    in_code = False
    code_lines = []
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            paragraph = doc.add_paragraph(" ".join(part.strip() for part in paragraph_buffer))
            paragraph.paragraph_format.widow_control = True
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_table(doc, table_lines)
            continue
        heading = re.match(r"^(#{2,4})\s+(.*)$", stripped)
        if heading:
            flush_paragraph()
            level = min(len(heading.group(1)) - 1, 3)
            text = clean_inline(heading.group(2))
            doc.add_heading(text, level=level)
            if text == "Purpose":
                pass
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet or numbered:
            flush_paragraph()
            text = clean_inline((bullet or numbered).group(1))
            if bullet:
                paragraph = doc.add_paragraph(text, style="List Bullet")
            else:
                number = re.match(r"^(\d+)\.", stripped).group(1)
                paragraph = doc.add_paragraph(f"{number}.  {text}")
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.space_after = Pt(2)
            index += 1
            continue
        paragraph_buffer.append(clean_inline(stripped))
        index += 1
    flush_paragraph()


def build():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    FLOWCHART.parent.mkdir(parents=True, exist_ok=True)
    create_flowchart()

    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Anaplan Daily Backup Project Documentation")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("Operator and maintenance guide")
    set_run_font(run, size=14, color=NAVY)

    intro = doc.add_paragraph(
        "This manual explains how to install, configure, run, schedule, monitor, "
        "and troubleshoot the Anaplan Daily Backup project. It covers single-model "
        "testing and sequential backups across multiple workspaces."
    )
    intro.paragraph_format.space_after = Pt(16)

    scope = doc.add_paragraph()
    label = scope.add_run("Primary outcome: ")
    set_run_font(label, bold=True, color=NAVY)
    scope.add_run(
        "each enabled model receives a verified timestamped Copy & Archive backup, "
        "with a JSON result report and browser cleanup at the end of the run."
    )

    doc.add_page_break()
    doc.add_heading("Backup Process Flow", level=1)
    paragraph = doc.add_paragraph(
        "The workflow separates dry runs from live browser operations. Live batches "
        "authenticate once, process enabled targets one at a time, record each result, "
        "and close the browser after the final target."
    )
    paragraph.paragraph_format.keep_with_next = True
    doc.add_picture(str(FLOWCHART), width=Inches(6.65))
    caption = doc.add_paragraph("Figure 1  Multi-model backup and archive workflow")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    for run in caption.runs:
        set_run_font(run, size=8.5, color="666666")
    doc.add_page_break()

    markdown = SOURCE.read_text(encoding="utf-8")
    add_markdown_content(doc, markdown)

    for section in doc.sections:
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.72)
        section.left_margin = Inches(0.82)
        section.right_margin = Inches(0.82)
        add_page_number(section.footer.paragraphs[0])

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.name is None:
                set_run_font(run)

    core = doc.core_properties
    core.title = "Anaplan Daily Backup Project Documentation"
    core.subject = "Installation, configuration, operation, scheduling, and troubleshooting"
    core.author = "Anaplan Daily Backup Project"
    core.keywords = "Anaplan, backup, archive, operations, Playwright"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
